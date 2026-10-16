from datetime import date, datetime, time, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import from_excel

from ..dependencies import require_admin
from ..supabase import db


router = APIRouter(
    prefix="/api/itineraries",
    tags=["Import Export Itinerary"],
)

MAX_FILE_SIZE = 10 * 1024 * 1024

ITINERARY_HEADERS = [
    "kode_itinerary",
    "judul",
    "tanggal_mulai",
    "tanggal_selesai",
]

ACTIVITY_HEADERS = [
    "kode_itinerary",
    "hari_ke",
    "tanggal",
    "judul_hari",
    "waktu",
    "kegiatan",
    "deskripsi",
    "lokasi",
    "urutan",
]

RECIPIENT_HEADERS = [
    "kode_itinerary",
    "role_penerima",
    "mode_penerima",
    "email_penerima",
]


# =====================================================
# EXCEL HELPER
# =====================================================

def style_workbook(workbook):
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column in ws.columns:
            values = [
                len(str(cell.value))
                for cell in column
                if cell.value is not None
            ]

            width = min(max(values, default=10) + 2, 45)
            ws.column_dimensions[
                column[0].column_letter
            ].width = max(width, 12)


def workbook_response(workbook, filename: str):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


def empty_to_none(value):
    if value is None:
        return None

    if isinstance(value, str) and not value.strip():
        return None

    return value


def text_value(value):
    value = empty_to_none(value)

    if value is None:
        return None

    return str(value).strip()


def parse_excel_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        converted = from_excel(value)

        if isinstance(converted, datetime):
            return converted.date()

        if isinstance(converted, date):
            return converted

    if isinstance(value, str):
        return date.fromisoformat(value.strip()[:10])

    raise ValueError("Tanggal Excel tidak valid.")


def parse_excel_time(value) -> time | None:
    value = empty_to_none(value)

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)

    if isinstance(value, time):
        return value.replace(tzinfo=None)

    if isinstance(value, (int, float)):
        seconds = round(float(value) * 86400) % 86400
        return (
            datetime.min + timedelta(seconds=seconds)
        ).time()

    if isinstance(value, str):
        return time.fromisoformat(value.strip())

    raise ValueError("Format waktu kegiatan tidak valid.")


def read_sheet(workbook, sheet_name: str, required_headers: list[str]):
    sheet_lookup = {
        name.strip().casefold(): name
        for name in workbook.sheetnames
    }

    actual_name = sheet_lookup.get(sheet_name.casefold())

    if not actual_name:
        raise ValueError(
            f"Sheet '{sheet_name}' tidak ditemukan."
        )

    ws = workbook[actual_name]
    headers = {}

    for index, cell in enumerate(ws[1]):
        name = text_value(cell.value)

        if name:
            headers[name.casefold()] = index

    missing = [
        header
        for header in required_headers
        if header.casefold() not in headers
    ]

    if missing:
        raise ValueError(
            f"Sheet '{sheet_name}' kekurangan kolom: "
            f"{', '.join(missing)}."
        )

    rows = []

    for row_number, values in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if all(empty_to_none(value) is None for value in values):
            continue

        row = {
            header: (
                values[index]
                if index < len(values)
                else None
            )
            for header, index in headers.items()
        }

        row["_row"] = row_number
        rows.append(row)

    return rows


def required_text(row: dict, column: str, sheet_name: str) -> str:
    value = text_value(row.get(column))

    if not value:
        raise ValueError(
            f"Sheet '{sheet_name}', baris {row['_row']}: "
            f"kolom '{column}' wajib diisi."
        )

    return value


def parse_recipient_role(value: str) -> str:
    normalized = value.strip().casefold().replace(" ", "_")

    if normalized not in {"tour_leader", "muthawif"}:
        raise ValueError(
            "Role penerima harus tour_leader atau muthawif."
        )

    return normalized


# =====================================================
# GET TEMPLATE EXCEL
# GET /api/itineraries/template
# =====================================================

@router.get("/template")
def download_itinerary_template(
    _=Depends(require_admin),
):
    wb = Workbook()

    instructions = wb.active
    instructions.title = "Petunjuk"

    instructions.append(["Bagian", "Petunjuk"])
    instructions.append([
        "Itinerary",
        "Satu baris untuk setiap rundown. Kode harus unik.",
    ])
    instructions.append([
        "Kegiatan",
        "Satu baris untuk setiap kegiatan. "
        "Hari tanpa kegiatan boleh memiliki baris kosong kegiatan.",
    ])
    instructions.append([
        "Penerima",
        "Gunakan all untuk semua role atau selected untuk akun tertentu.",
    ])
    instructions.append([
        "Kode",
        "kode_itinerary harus konsisten pada semua sheet.",
    ])
    instructions.append([
        "Tanggal",
        "Gunakan format tanggal Excel atau YYYY-MM-DD.",
    ])
    instructions.append([
        "Waktu",
        "Gunakan format HH:MM atau HH:MM:SS.",
    ])
    instructions.append([
        "Email",
        "Untuk selected, gunakan email yang terdaftar di profiles.",
    ])
    instructions.append([
        "Status",
        "Hasil import selalu draft untuk diperiksa admin.",
    ])

    itineraries = wb.create_sheet("Itinerary")
    itineraries.append(ITINERARY_HEADERS)

    activities = wb.create_sheet("Kegiatan")
    activities.append(ACTIVITY_HEADERS)

    recipients = wb.create_sheet("Penerima")
    recipients.append(RECIPIENT_HEADERS)

    style_workbook(wb)

    return workbook_response(
        wb,
        "template_manajemen_itinerary.xlsx",
    )


# =====================================================
# POST IMPORT EXCEL
# POST /api/itineraries/import
# =====================================================

@router.post("/import")
def import_itineraries(
    file: UploadFile = File(...),
    current=Depends(require_admin),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="File harus menggunakan format .xlsx.",
        )

    content = file.file.read(MAX_FILE_SIZE + 1)

    if not content:
        raise HTTPException(
            status_code=400,
            detail="File Excel kosong.",
        )

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="Ukuran file maksimal 10 MB.",
        )

    try:
        wb = load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=True,
        )

        itinerary_rows = read_sheet(
            wb,
            "Itinerary",
            [
                "kode_itinerary",
                "judul",
                "tanggal_mulai",
                "tanggal_selesai",
            ],
        )

        activity_rows = read_sheet(
            wb,
            "Kegiatan",
            [
                "kode_itinerary",
                "hari_ke",
                "tanggal",
                "kegiatan",
            ],
        )

        recipient_rows = read_sheet(
            wb,
            "Penerima",
            [
                "kode_itinerary",
                "role_penerima",
                "mode_penerima",
            ],
        )

        if not itinerary_rows:
            raise ValueError(
                "Sheet Itinerary belum berisi data."
            )

        # ---------------------------------------------
        # 1. Validasi data utama itinerary
        # ---------------------------------------------

        itinerary_data = {}

        for row in itinerary_rows:
            code = required_text(
                row,
                "kode_itinerary",
                "Itinerary",
            )

            if code in itinerary_data:
                raise ValueError(
                    f"Kode itinerary '{code}' duplikat."
                )

            title = required_text(
                row,
                "judul",
                "Itinerary",
            ).strip()

            start_date = parse_excel_date(
                row.get("tanggal_mulai")
            )

            end_date = parse_excel_date(
                row.get("tanggal_selesai")
            )

            if end_date < start_date:
                raise ValueError(
                    f"Periode itinerary '{code}' tidak valid."
                )

            itinerary_data[code] = {
                "title": title,
                "start_date": start_date,
                "end_date": end_date,
            }

        # ---------------------------------------------
        # 2. Validasi kegiatan
        # ---------------------------------------------

        activity_data = {code: [] for code in itinerary_data}
        day_titles = {}
        used_orders = set()
        next_orders = {}

        for row in activity_rows:
            code = required_text(
                row,
                "kode_itinerary",
                "Kegiatan",
            )

            if code not in itinerary_data:
                raise ValueError(
                    f"Kode '{code}' di sheet Kegiatan "
                    "tidak ditemukan di sheet Itinerary."
                )

            try:
                day_number = int(row.get("hari_ke"))

                if day_number < 1:
                    raise ValueError()

                activity_date = parse_excel_date(
                    row.get("tanggal")
                )

            except (TypeError, ValueError):
                raise ValueError(
                    f"Sheet Kegiatan, baris {row['_row']}: "
                    "hari_ke atau tanggal tidak valid."
                )

            itinerary = itinerary_data[code]

            if not (
                itinerary["start_date"]
                <= activity_date
                <= itinerary["end_date"]
            ):
                raise ValueError(
                    f"Tanggal kegiatan di baris {row['_row']} "
                    "berada di luar periode itinerary."
                )

            expected_day = (
                activity_date - itinerary["start_date"]
            ).days + 1

            if day_number != expected_day:
                raise ValueError(
                    f"Sheet Kegiatan, baris {row['_row']}: "
                    f"tanggal tersebut harus menjadi hari ke-{expected_day}."
                )

            day_key = (code, day_number)
            day_title = text_value(row.get("judul_hari"))

            if day_title:
                existing_title = day_titles.get(day_key)

                if existing_title and existing_title != day_title:
                    raise ValueError(
                        f"Judul hari untuk {code}, hari {day_number} "
                        "tidak konsisten."
                    )

                day_titles[day_key] = day_title

            activity_title = text_value(row.get("kegiatan"))

            # Baris tanpa nama kegiatan masih dapat dipakai
            # untuk menyimpan judul hari.
            if not activity_title:
                continue

            try:
                activity_time = parse_excel_time(
                    row.get("waktu")
                )
            except ValueError as exc:
                raise ValueError(
                    f"Sheet Kegiatan, baris {row['_row']}: {exc}"
                )

            raw_order = empty_to_none(row.get("urutan"))

            if raw_order is None:
                sort_order = next_orders.get(day_key, 0) + 1
            else:
                try:
                    sort_order = int(raw_order)

                    if sort_order < 1:
                        raise ValueError()

                except (TypeError, ValueError):
                    raise ValueError(
                        f"Urutan kegiatan pada baris "
                        f"{row['_row']} tidak valid."
                    )

            order_key = (code, day_number, sort_order)

            if order_key in used_orders:
                raise ValueError(
                    f"Urutan kegiatan duplikat pada itinerary "
                    f"'{code}', hari ke-{day_number}."
                )

            used_orders.add(order_key)
            next_orders[day_key] = max(
                next_orders.get(day_key, 0),
                sort_order,
            )

            activity_data[code].append({
                "day_number": day_number,
                "activity_date": activity_date,
                "activity_time": activity_time,
                "title": activity_title,
                "description": text_value(
                    row.get("deskripsi")
                ),
                "location": text_value(
                    row.get("lokasi")
                ),
                "sort_order": sort_order,
            })

        for code, activities in activity_data.items():
            if not activities:
                raise ValueError(
                    f"Itinerary '{code}' belum memiliki kegiatan."
                )

        # ---------------------------------------------
        # 3. Validasi aturan penerima
        # ---------------------------------------------

        recipient_data = {
            code: {}
            for code in itinerary_data
        }

        for row in recipient_rows:
            code = required_text(
                row,
                "kode_itinerary",
                "Penerima",
            )

            if code not in itinerary_data:
                raise ValueError(
                    f"Kode '{code}' di sheet Penerima "
                    "tidak ditemukan di sheet Itinerary."
                )

            role = parse_recipient_role(
                required_text(
                    row,
                    "role_penerima",
                    "Penerima",
                )
            )

            mode = required_text(
                row,
                "mode_penerima",
                "Penerima",
            ).casefold()

            if mode not in {"all", "selected"}:
                raise ValueError(
                    "mode_penerima harus all atau selected."
                )

            email = text_value(row.get("email_penerima"))

            existing = recipient_data[code].get(role)

            if existing and existing["mode"] != mode:
                raise ValueError(
                    f"Mode penerima untuk {code}, {role} "
                    "tidak boleh berbeda."
                )

            if mode == "all":
                if email:
                    raise ValueError(
                        "Email penerima harus kosong jika mode_penerima=all."
                    )

                if existing:
                    raise ValueError(
                        f"Aturan penerima all duplikat untuk {code}, {role}."
                    )

                recipient_data[code][role] = {
                    "mode": "all",
                    "emails": [],
                }

            else:
                if not email or "@" not in email:
                    raise ValueError(
                        f"Email wajib diisi untuk penerima selected "
                        f"di itinerary '{code}'."
                    )

                if existing is None:
                    recipient_data[code][role] = {
                        "mode": "selected",
                        "emails": [],
                    }

                emails = recipient_data[code][role]["emails"]
                normalized_email = email.casefold()

                if normalized_email in emails:
                    raise ValueError(
                        f"Email penerima {email} duplikat "
                        f"untuk itinerary '{code}'."
                    )

                emails.append(normalized_email)

        for code, rules in recipient_data.items():
            if not rules:
                raise ValueError(
                    f"Itinerary '{code}' belum memiliki aturan penerima."
                )

            for role, rule in rules.items():
                if (
                    rule["mode"] == "selected"
                    and not rule["emails"]
                ):
                    raise ValueError(
                        f"Penerima selected belum dipilih "
                        f"untuk itinerary '{code}', role {role}."
                    )

    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise

        raise HTTPException(
            status_code=422,
            detail=f"Format atau isi Excel tidak valid: {exc}",
        )

    # ---------------------------------------------
    # 4. Validasi profil penerima dari Supabase
    # Dilakukan sebelum mulai menyimpan itinerary.
    # ---------------------------------------------

    selected_emails = {
        email
        for rules in recipient_data.values()
        for rule in rules.values()
        if rule["mode"] == "selected"
        for email in rule["emails"]
    }

    profiles_by_email = {}

    if selected_emails:
        profiles = (
            db.table("profiles")
            .select("id, email, role, is_active")
            .execute()
            .data or []
        )

        profiles_by_email = {
            item["email"].strip().casefold(): item
            for item in profiles
            if item.get("email")
        }

        for code, rules in recipient_data.items():
            for role, rule in rules.items():
                if rule["mode"] != "selected":
                    continue

                for email in rule["emails"]:
                    profile = profiles_by_email.get(email)

                    if not profile:
                        raise HTTPException(
                            status_code=422,
                            detail=(
                                f"Email penerima '{email}' pada "
                                f"itinerary '{code}' tidak ditemukan."
                            ),
                        )

                    if (
                        profile["role"] != role
                        or not profile["is_active"]
                    ):
                        raise HTTPException(
                            status_code=422,
                            detail=(
                                f"Role akun '{email}' tidak cocok "
                                "atau akun tidak aktif."
                            ),
                        )

    # ---------------------------------------------
    # 5. Simpan itinerary dan seluruh data anak
    # ---------------------------------------------

    created_ids = []

    try:
        for code, itinerary in itinerary_data.items():
            parent_result = (
                db.table("itineraries")
                .insert({
                    "title": itinerary["title"],
                    "start_date": itinerary[
                        "start_date"
                    ].isoformat(),
                    "end_date": itinerary[
                        "end_date"
                    ].isoformat(),
                    "status": "draft",
                    "created_by": current["profile"]["id"],
                })
                .execute()
            )

            parents = parent_result.data or []

            if not parents:
                raise RuntimeError(
                    f"Itinerary '{code}' gagal disimpan."
                )

            itinerary_id = parents[0]["id"]
            created_ids.append(itinerary_id)

            # Buat seluruh hari secara otomatis.
            total_days = (
                itinerary["end_date"]
                - itinerary["start_date"]
            ).days + 1

            day_records = []

            for offset in range(total_days):
                activity_date = (
                    itinerary["start_date"]
                    + timedelta(days=offset)
                )

                day_number = offset + 1

                day_records.append({
                    "itinerary_id": itinerary_id,
                    "day_number": day_number,
                    "activity_date": activity_date.isoformat(),
                    "title": day_titles.get(
                        (code, day_number)
                    ),
                })

            (
                db.table("itinerary_days")
                .insert(day_records)
                .execute()
            )

            saved_days = (
                db.table("itinerary_days")
                .select("id, day_number, activity_date")
                .eq("itinerary_id", itinerary_id)
                .execute()
                .data or []
            )

            day_ids = {
                day["day_number"]: day["id"]
                for day in saved_days
            }

            # Simpan seluruh kegiatan berdasarkan hari.
            records = []

            for activity in activity_data[code]:
                day_number = activity["day_number"]

                records.append({
                    "itinerary_day_id": day_ids[day_number],
                    "activity_time": (
                        activity["activity_time"].isoformat()
                        if activity["activity_time"]
                        else None
                    ),
                    "title": activity["title"],
                    "description": activity["description"],
                    "location": activity["location"],
                    "sort_order": activity["sort_order"],
                })

            if records:
                (
                    db.table("itinerary_activities")
                    .insert(records)
                    .execute()
                )

            # Simpan aturan penerima.
            for role, rule in recipient_data[code].items():
                rule_result = (
                    db.table("itinerary_recipient_rules")
                    .insert({
                        "itinerary_id": itinerary_id,
                        "recipient_role": role,
                        "recipient_mode": rule["mode"],
                    })
                    .execute()
                )

                saved_rules = rule_result.data or []

                if not saved_rules:
                    raise RuntimeError(
                        "Aturan penerima gagal disimpan."
                    )

                rule_id = saved_rules[0]["id"]

                if rule["mode"] == "selected":
                    recipient_records = [
                        {
                            "rule_id": rule_id,
                            "profile_id": profiles_by_email[
                                email
                            ]["id"],
                        }
                        for email in rule["emails"]
                    ]

                    (
                        db.table("itinerary_recipient_users")
                        .insert(recipient_records)
                        .execute()
                    )

    except Exception as exc:
        # Best-effort cleanup: hapus seluruh itinerary yang
        # berhasil dibuat. Foreign key CASCADE menghapus data anak.
        if created_ids:
            try:
                (
                    db.table("itineraries")
                    .delete()
                    .in_("id", created_ids)
                    .execute()
                )
            except Exception:
                pass

        raise HTTPException(
            status_code=500,
            detail=(
                "Import gagal saat menyimpan data. "
                f"Detail: {exc}"
            ),
        )

    return {
        "message": "Import itinerary berhasil.",
        "imported_count": len(created_ids),
        "data": [
            {
                "id": itinerary_id,
                "title": itinerary_data[code]["title"],
                "status": "draft",
            }
            for code, itinerary_id in zip(
                itinerary_data.keys(),
                created_ids,
            )
        ],
    }


# =====================================================
# GET EXPORT EXCEL
# GET /api/itineraries/export
# =====================================================

@router.get("/export")
def export_itineraries(
    status: str | None = None,
    _=Depends(require_admin),
):
    if status is not None and status not in {
        "draft",
        "published",
        "archived",
    }:
        raise HTTPException(
            status_code=400,
            detail="Status itinerary tidak valid.",
        )

    query = (
        db.table("itineraries")
        .select("*")
        .order("created_at", desc=True)
    )

    if status:
        query = query.eq("status", status)

    itineraries = query.execute().data or []

    wb = Workbook()

    ws_itinerary = wb.active
    ws_itinerary.title = "Itinerary"
    ws_itinerary.append([
        "kode_itinerary",
        "judul",
        "tanggal_mulai",
        "tanggal_selesai",
        "status",
        "dibuat",
    ])

    ws_activities = wb.create_sheet("Kegiatan")
    ws_activities.append(ACTIVITY_HEADERS)

    ws_recipients = wb.create_sheet("Penerima")
    ws_recipients.append(RECIPIENT_HEADERS)

    for itinerary in itineraries:
        itinerary_id = itinerary["id"]
        code = itinerary_id

        ws_itinerary.append([
            code,
            itinerary["title"],
            itinerary["start_date"],
            itinerary["end_date"],
            itinerary["status"],
            itinerary["created_at"],
        ])

        days = (
            db.table("itinerary_days")
            .select("*")
            .eq("itinerary_id", itinerary_id)
            .order("day_number")
            .execute()
            .data or []
        )

        for day in days:
            activities = (
                db.table("itinerary_activities")
                .select("*")
                .eq("itinerary_day_id", day["id"])
                .order("sort_order")
                .execute()
                .data or []
            )

            if not activities:
                ws_activities.append([
                    code,
                    day["day_number"],
                    str(day["activity_date"])[:10],
                    day.get("title"),
                    None,
                    None,
                    None,
                    None,
                    None,
                ])

            for activity in activities:
                ws_activities.append([
                    code,
                    day["day_number"],
                    str(day["activity_date"])[:10],
                    day.get("title"),
                    activity.get("activity_time"),
                    activity["title"],
                    activity.get("description"),
                    activity.get("location"),
                    activity["sort_order"],
                ])

        rules = (
            db.table("itinerary_recipient_rules")
            .select("*")
            .eq("itinerary_id", itinerary_id)
            .execute()
            .data or []
        )

        for rule in rules:
            if rule["recipient_mode"] == "all":
                ws_recipients.append([
                    code,
                    rule["recipient_role"],
                    "all",
                    None,
                ])
                continue

            links = (
                db.table("itinerary_recipient_users")
                .select("profile_id")
                .eq("rule_id", rule["id"])
                .execute()
                .data or []
            )

            profile_ids = [
                item["profile_id"] for item in links
            ]

            profiles = []

            if profile_ids:
                profiles = (
                    db.table("profiles")
                    .select("email")
                    .in_("id", profile_ids)
                    .execute()
                    .data or []
                )

            if not profiles:
                ws_recipients.append([
                    code,
                    rule["recipient_role"],
                    "selected",
                    None,
                ])

            for profile in profiles:
                ws_recipients.append([
                    code,
                    rule["recipient_role"],
                    "selected",
                    profile.get("email"),
                ])

    style_workbook(wb)

    return workbook_response(
        wb,
        "manajemen_itinerary.xlsx",
    )